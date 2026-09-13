"""Azure OpenAI (GPT) を使用した PowerPoint ネイティブプレゼンテーション自動生成モジュール。"""
# pylint: disable=too-many-lines,line-too-long,too-many-locals,too-many-arguments,too-many-positional-arguments,too-few-public-methods,broad-exception-caught,unused-argument
from __future__ import annotations

import base64
import io
import json
import os
import shutil
import traceback
from typing import Any, Dict, List, Optional

from PIL import Image
import streamlit as st
from pptx import Presentation

try:
    from gp_chat import azure_history_utils
    from gp_chat import state_manager
    from gp_chat import utils
    from gp_chat.azure_responses_router import _build_client, generate_response
    from gp_chat.azure_runtime import AzureRuntime
    from gp_chat.pptx_agent import (
        CropAreaSchema,
        PlaceholderContent,
        PresentationDSLSchema,
        PresentationSourceBrief,
        SlideNode,
        _attach_reference_usage,
        _brief_to_text,
        _build_conversation_excerpt,
        _finalize_reference_entries,
        _format_attachment_summary,
        _minimum_body_slide_count,
        generate_slide_html,
        save_physical_presentation,
        scan_template_layouts,
        validate_slide_overflow,
    )
except ImportError:
    import azure_history_utils
    import state_manager
    import utils
    from azure_responses_router import _build_client, generate_response
    from azure_runtime import AzureRuntime
    from pptx_agent import (
        CropAreaSchema,
        PlaceholderContent,
        PresentationDSLSchema,
        PresentationSourceBrief,
        SlideNode,
        _attach_reference_usage,
        _brief_to_text,
        _build_conversation_excerpt,
        _finalize_reference_entries,
        _format_attachment_summary,
        _minimum_body_slide_count,
        generate_slide_html,
        save_physical_presentation,
        scan_template_layouts,
        validate_slide_overflow,
    )


def _safe_json_loads(raw_text: str) -> dict[str, Any]:
    """Markdownコードブロックを取り除き安全にJSONをパースする。"""
    clean_text = (raw_text or "").strip()
    if clean_text.startswith("```"):
        lines = clean_text.split("\n")
        if len(lines) >= 3:
            clean_text = "\n".join(lines[1:-1]).strip()
        else:
            clean_text = clean_text.replace("```json", "").replace("```", "").strip()
    start_idx = clean_text.find("{")
    end_idx = clean_text.rfind("}")
    if start_idx != -1 and end_idx != -1 and end_idx >= start_idx:
        clean_text = clean_text[start_idx : end_idx + 1]
    return json.loads(clean_text)


def _parse_source_brief(raw_text: str) -> PresentationSourceBrief:
    """GPTの応答テキストからPresentationSourceBriefをパースする。"""
    data = _safe_json_loads(raw_text)
    return PresentationSourceBrief.model_validate(data)


def _parse_presentation_dsl(raw_text: str) -> PresentationDSLSchema:
    """GPTの応答テキストからPresentationDSLSchemaをパースする。"""
    data = _safe_json_loads(raw_text)
    return PresentationDSLSchema.model_validate(data)


def _generate_dalle_image(runtime: AzureRuntime, prompt: str, output_path: str) -> bool:
    """Azure OpenAI の DALL-E 3 を呼び出してスライド用イラストを生成。未設定・失敗時は安全にスキップ。"""
    dalle_deployment = getattr(runtime, "dalle_deployment", "") or ""
    if not dalle_deployment:
        state_manager.add_debug_log(
            "[AzurePPTX] DALL-E deployment not configured (AZURE_OPENAI_DALLE_DEPLOYMENT is empty). Skipping AI image generation.",
            "info",
        )
        return False

    state_manager.add_debug_log(
        f"[AzurePPTX] Requesting DALL-E image (deployment={dalle_deployment}). Prompt: {prompt[:100]}..."
    )
    try:
        client = _build_client(runtime)
        response = client.images.generate(
            model=dalle_deployment,
            prompt=prompt[:1000],
            n=1,
            size="1024x1024",
            response_format="b64_json",
        )
        if response and response.data and response.data[0].b64_json:
            img_bytes = base64.b64decode(response.data[0].b64_json)
            with open(output_path, "wb") as f:
                f.write(img_bytes)
            state_manager.add_debug_log(f"[AzurePPTX] DALL-E image saved to {output_path}")
            return True
        state_manager.add_debug_log("[AzurePPTX] DALL-E returned empty image data.", "warning")
        return False
    except Exception as e:
        state_manager.add_debug_log(f"[AzurePPTX] DALL-E image generation failed: {e}", "warning")
        return False


def _crop_user_image_with_azure(
    runtime: AzureRuntime,
    image_bytes: bytes,
    image_mime_type: str,
    instruction: str,
    output_path: str,
) -> bool:
    """GPT-4o/5 Vision を用いてユーザー添付画像のバウンディングボックスを検知し、Pillowでトリミング。"""
    state_manager.add_debug_log(f"[AzurePPTX] Cropping user image based on instruction: {instruction}")
    try:
        b64_data = base64.b64encode(image_bytes).decode("ascii")
        prompt_text = (
            f"ユーザー指示: 「{instruction}」\n"
            "指示された領域を特定し、そのバウンディングボックスの相対座標（0-1000の整数）をJSONで返してください。\n"
            "JSONスキーマ: {\"ymin\": 0-1000, \"xmin\": 0-1000, \"ymax\": 0-1000, \"xmax\": 0-1000}"
        )
        messages = [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt_text},
                    {"type": "image_url", "image_url": {"url": f"data:{image_mime_type};base64,{b64_data}"}},
                ],
            }
        ]
        result = generate_response(
            runtime=runtime,
            input_messages=messages,
            instructions="あなたは画像解析の専門家です。指定されたJSONスキーマに厳密に従ってバウンディングボックスを返してください。",
            max_output_tokens=1024,
            temperature=0.1,
            response_mime_type="application/json",
        )
        crop_data = _safe_json_loads(result.text)
        coords = CropAreaSchema.model_validate(crop_data)

        orig_img = Image.open(io.BytesIO(image_bytes))
        width, height = orig_img.size
        box = (
            int(coords.xmin * width / 1000.0),
            int(coords.ymin * height / 1000.0),
            int(coords.xmax * width / 1000.0),
            int(coords.ymax * height / 1000.0),
        )
        cropped = orig_img.crop(box)
        cropped.save(output_path)
        state_manager.add_debug_log(f"[AzurePPTX] Cropped user image saved to {output_path}")
        return True
    except Exception as e:
        state_manager.add_debug_log(f"[AzurePPTX] User image cropping failed: {e}. Saving raw image.", "warning")
        try:
            with open(output_path, "wb") as f:
                f.write(image_bytes)
            return True
        except Exception:
            return False


class AzurePPTXAgent:
    """Azure OpenAI (GPT) を使用した PowerPoint ネイティブスライド自動生成エージェント。"""

    def __init__(self, runtime: AzureRuntime):
        self.runtime = runtime
        state_manager.add_debug_log("[AzurePPTXAgent] Initialized successfully.")

    def _run_research_context(
        self,
        context_messages: list[dict[str, Any]],
        system_instruction: str,
        search_enabled: bool,
    ) -> tuple[str, dict | None]:
        if not search_enabled:
            return "", None
        try:
            research_messages = list(context_messages)
            research_messages.append(
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "input_text",
                            "text": (
                                "PowerPoint資料化の前提調査として、会話の目的に必要な最新情報・事実関係・"
                                "重要な出典候補を簡潔に整理してください。"
                                "スライドに入れるべき具体的な事実・数値・日付・比較軸を抽出してください。"
                            ),
                        }
                    ],
                }
            )
            result = generate_response(
                runtime=self.runtime,
                input_messages=research_messages,
                instructions=(
                    "あなたは資料作成前のリサーチ担当です。最新情報が必要な論点はWeb検索を行い、"
                    "スライドに使える事実、日付、比較軸、注意点を日本語で簡潔に整理してください。"
                ),
                max_output_tokens=4096,
                temperature=0.1,
                search_enabled=True,
            )
            return result.text or "", result.grounding_metadata
        except Exception as e:
            state_manager.add_debug_log(f"[AzurePPTXAgent] Web research step skipped: {e}", "warning")
            return "", None

    def _build_source_brief(
        self,
        context_messages: list[dict[str, Any]],
        conversation_excerpt: str,
        attachment_summary: str,
        file_attachments_meta: Optional[List[dict]],
        research_context: str,
        grounding_metadata: Optional[dict],
        system_instruction: str,
    ) -> PresentationSourceBrief:
        prompt = (
            "これまでの会話、添付ファイル、検索結果を統合し、PowerPoint資料を作るための材料台帳 (source brief JSON) を作成してください。\n"
            "事実と推測を分け、スライド構成に必要な数値・日付・比較軸・出典候補を漏れなく整理してください。\n\n"
            + (f"【会話ログ抜粋】\n{conversation_excerpt}\n\n" if conversation_excerpt else "")
            + (f"【添付ファイル情報】\n{attachment_summary}\n\n" if attachment_summary else "")
            + (f"【Web検索リサーチ結果】\n{research_context}\n\n" if research_context else "")
            + (
                f"【検索メタデータ】\n{json.dumps(grounding_metadata, ensure_ascii=False, indent=2)}\n\n"
                if grounding_metadata
                else ""
            )
        )
        messages = list(context_messages)
        messages.append({"role": "user", "content": [{"type": "input_text", "text": prompt}]})

        state_manager.add_debug_log("[AzurePPTXAgent] Requesting source brief JSON from GPT...")
        result = generate_response(
            runtime=self.runtime,
            input_messages=messages,
            instructions=(
                "あなたは資料作成前の編集長です。会話、添付、検索結果を統合し、"
                "スライド構成に使える材料台帳を完全なJSONで作成してください。"
            ),
            max_output_tokens=16384,
            temperature=0.1,
            response_mime_type="application/json",
            structured_output_name="source_brief",
        )
        brief = _parse_source_brief(result.text)
        state_manager.add_debug_log(
            f"[AzurePPTXAgent] Source brief built: facts={len(brief.key_facts)}, "
            f"units={len(brief.source_coverage_units)}."
        )

        if not brief.source_coverage_units:
            fallback_units = []
            fallback_units.extend(brief.coverage_requirements[:20])
            fallback_units.extend(brief.key_facts[:30])
            brief.source_coverage_units = list(dict.fromkeys(unit for unit in fallback_units if unit))

        brief.references = _finalize_reference_entries(
            references=brief.references,
            file_attachments_meta=file_attachments_meta,
            grounding_metadata=grounding_metadata,
        )
        return brief

    def _generate_presentation_structure(
        self,
        brief: PresentationSourceBrief,
        layouts_info: Dict[str, Any],
        has_template: bool,
    ) -> PresentationDSLSchema:
        min_body_slides = _minimum_body_slide_count(brief)
        template_instruction = ""
        if has_template and layouts_info:
            allowed_layouts = ", ".join(f"'{name}'" for name in layouts_info.keys())
            template_instruction = (
                "\n\n【利用可能なスライドレイアウト情報】\n"
                f"各スライドの 'layout_name' は必ず次の候補から選択してください: {allowed_layouts}\n"
                "表紙用・裏表紙用のレイアウトは本文スライドに使用しないでください。\n"
                "'placeholders' では、指定したレイアウトに定義された 'idx' だけを指定してください。\n"
            )

        prompt = (
            "以下の source brief を唯一の材料台帳として、最高品質のスライド構成案JSONを出力してください。\n"
            f"本文スライドは最低 {min_body_slides} 枚作成してください。\n"
            "各スライドにはタイトル、本文プレースホルダーへのテキスト（箇条書きは改行区切り）、"
            "適切なビジュアル要素（visual_type: timeline/process/comparison/kpi/matrix/risk/summary/auto）を指定してください。\n"
            + template_instruction
            + f"\n\n【source brief】\n{_brief_to_text(brief)}"
        )
        messages = [{"role": "user", "content": [{"type": "input_text", "text": prompt}]}]

        state_manager.add_debug_log("[AzurePPTXAgent] Requesting presentation structure JSON from GPT...")
        result = generate_response(
            runtime=self.runtime,
            input_messages=messages,
            instructions="あなたは一流のプレゼンテーションデザイナーです。Pydantic PresentationDSLSchema に完全に適合するJSONを出力してください。",
            max_output_tokens=16384,
            temperature=0.2,
            response_mime_type="application/json",
            structured_output_name="presentation_dsl",
        )
        dsl = _parse_presentation_dsl(result.text)
        _attach_reference_usage(brief, dsl)
        state_manager.add_debug_log(
            f"[AzurePPTXAgent] Structure generated: title='{dsl.presentation_title}', slides={len(dsl.slides)}"
        )
        return dsl

    def _validate_and_adjust_slides(
        self,
        slides: List[SlideNode],
        layouts_info: Dict[str, Any],
        temp_dir: str,
    ) -> tuple[List[SlideNode], List[int]]:
        offsets = [0] * len(slides)
        for retry_loop in range(3):
            state_manager.add_debug_log(f"[AzurePPTXAgent] Geometry Validation Round {retry_loop + 1}")
            for i, slide in enumerate(slides):
                html_content = generate_slide_html(slide, layouts_info, font_size_offset=offsets[i])
                html_path = os.path.join(temp_dir, f"{slide.slide_number:02d}.html")
                with open(html_path, "w", encoding="utf-8") as f:
                    f.write(html_content)

            validation_results = validate_slide_overflow(temp_dir, slides, offsets)
            overflowed_slides = [res for res in validation_results if res.get("overflowed")]
            if not overflowed_slides:
                state_manager.add_debug_log("[AzurePPTXAgent] Geometry validation passed. No overflows detected!")
                break

            if retry_loop == 2:
                state_manager.add_debug_log(
                    "[AzurePPTXAgent] Retries exhausted. Applying emergency font shrink (-2pt) for overflowed slides."
                )
                for res in overflowed_slides:
                    idx = next(
                        (i for i, s in enumerate(slides) if s.slide_number == res["slide_number"]),
                        None,
                    )
                    if idx is not None:
                        offsets[idx] = -2
                break

            state_manager.add_debug_log(
                f"[AzurePPTXAgent] Overflows detected on {len(overflowed_slides)} slides. Starting GPT self-correction..."
            )
            for res in overflowed_slides:
                slide_num = res["slide_number"]
                idx = next((i for i, s in enumerate(slides) if s.slide_number == slide_num), None)
                if idx is None:
                    continue
                target_slide = slides[idx]
                target_ph_idx = int(res["element_id"].split("_")[1])
                target_ph = next((p for p in target_slide.placeholders if p.idx == target_ph_idx), None)
                target_text = target_ph.text_content if target_ph else ""

                scale = res["available_height"] / max(res["required_height"], 1)
                retry_prompt = (
                    f"【レイアウト文字溢れ検知】\n"
                    f"スライド番号: {slide_num}\n"
                    f"エラー要素 (プレースホルダー idx): {target_ph_idx}\n"
                    f"現在のテキスト: \"{target_text}\"\n"
                    f"許容高さ: {res['available_height']}px / 要求高さ: {res['required_height']}px\n\n"
                    "上記要素で文字溢れが発生しました。"
                    f"スライド「{target_slide.title}」のプレースホルダー(idx={target_ph_idx})のテキストを、"
                    f"約 {scale:.2f} 倍（文字数を30%〜50%削減）に要約・短縮してください。\n"
                    "JSONスキーマ: {\"idx\": " + str(target_ph_idx) + ", \"text_content\": \"要約後のテキスト\"}"
                )
                try:
                    corr_result = generate_response(
                        runtime=self.runtime,
                        input_messages=[{"role": "user", "content": [{"type": "input_text", "text": retry_prompt}]}],
                        instructions="あなたは要約・短縮の専門家です。指定されたJSON形式で要約後のプレースホルダーを返してください。",
                        max_output_tokens=2048,
                        temperature=0.2,
                        response_mime_type="application/json",
                    )
                    corr_data = _safe_json_loads(corr_result.text)
                    new_content = PlaceholderContent.model_validate(corr_data)
                    for p_idx, p_item in enumerate(slides[idx].placeholders):
                        if p_item.idx == target_ph_idx:
                            slides[idx].placeholders[p_idx] = new_content
                            break
                    state_manager.add_debug_log(
                        f"[AzurePPTXAgent] Slide {slide_num} placeholder {target_ph_idx} text shortened successfully."
                    )
                except Exception as e:
                    state_manager.add_debug_log(f"[AzurePPTXAgent] Correction failed for slide {slide_num}: {e}", "warning")

        return slides, offsets

    def _prepare_slide_images(
        self,
        slides: List[SlideNode],
        temp_dir: str,
        user_images: Optional[List[dict]],
    ) -> Dict[int, Dict[int, str]]:
        slide_images: Dict[int, Dict[int, str]] = {}
        for slide in slides:
            slide_num = slide.slide_number
            slide_images[slide_num] = {}
            for content in slide.placeholders:
                if content.image_prompt:
                    out_path = os.path.join(temp_dir, f"ai_gen_{slide_num}_{content.idx}.png")
                    success = _generate_dalle_image(self.runtime, content.image_prompt, out_path)
                    if success:
                        slide_images[slide_num][content.idx] = out_path
                elif content.use_user_image and user_images:
                    user_img = user_images[0]
                    img_bytes = user_img["bytes"]
                    img_mime = user_img.get("type", "image/png")
                    out_path = os.path.join(temp_dir, f"user_img_{slide_num}_{content.idx}.png")
                    if content.crop_instruction:
                        success = _crop_user_image_with_azure(
                            self.runtime,
                            img_bytes,
                            img_mime,
                            content.crop_instruction,
                            out_path,
                        )
                        if success:
                            slide_images[slide_num][content.idx] = out_path
                    else:
                        try:
                            with open(out_path, "wb") as f:
                                f.write(img_bytes)
                            slide_images[slide_num][content.idx] = out_path
                        except Exception as e:
                            state_manager.add_debug_log(f"[AzurePPTXAgent] Saving user image failed: {e}", "warning")
        return slide_images

    def generate_presentation_pipeline(
        self,
        *,
        chat_history: List[dict],
        context,
        session_id: str,
        user_images: Optional[List[dict]] = None,
        search_enabled: bool = False,
        file_attachments_meta: Optional[List[dict]] = None,
    ) -> str:
        """Azure OpenAI を用いた PowerPoint 生成パイプライン本体。"""
        # 保存先ディレクトリの決定
        current_report_folder = st.session_state.get("current_report_folder")
        if not current_report_folder:
            current_chat_filename = st.session_state.get("current_chat_filename")
            if current_chat_filename:
                folder_name = os.path.splitext(os.path.basename(current_chat_filename))[0]
            else:
                folder_name = azure_history_utils.generate_chat_title(chat_history, self.runtime)
            folder_name = azure_history_utils.sanitize_filename(folder_name or "pptx_chat")[:80] or "pptx_chat"
            st.session_state["current_report_folder"] = folder_name
        else:
            folder_name = current_report_folder

        output_dir = os.path.join("slide_data", folder_name)
        os.makedirs(output_dir, exist_ok=True)
        temp_dir = os.path.join("temp_workspace", f"azure_{session_id[:8]}")
        os.makedirs(temp_dir, exist_ok=True)

        pptx_number = 1
        for filename in os.listdir(output_dir):
            stem, ext = os.path.splitext(filename)
            if ext.lower() == ".pptx" and stem.isdigit():
                pptx_number = max(pptx_number, int(stem) + 1)
        final_pptx_path = os.path.abspath(os.path.join(output_dir, f"{pptx_number:02d}.pptx"))

        # テンプレートのロードと解析
        current_dir = os.path.dirname(os.path.abspath(__file__))
        template_path = os.path.join(current_dir, "format.pptx")
        if not os.path.exists(template_path):
            template_path = os.path.join(current_dir, "format.potx")
        has_template = False
        layouts_info = {}
        check_prs = None
        if os.path.exists(template_path):
            try:
                check_prs = Presentation(template_path)
                if len(check_prs.slides) >= 3:
                    layouts_info = scan_template_layouts(template_path)
                    has_template = True
                    state_manager.add_debug_log(f"[AzurePPTXAgent] Scanned {len(layouts_info)} layouts from template.")
            except Exception as e:
                state_manager.add_debug_log(f"[AzurePPTXAgent] Template inspection warning: {e}", "warning")

        # 会話抜粋と添付情報の作成
        conversation_excerpt = _build_conversation_excerpt(chat_history)
        attachment_summary = _format_attachment_summary(file_attachments_meta)

        # 1. 前提リサーチ（Web検索が有効な場合）
        research_context, grounding_metadata = self._run_research_context(
            context_messages=context.messages,
            system_instruction=context.system_instruction,
            search_enabled=search_enabled,
        )

        # 2. Source Brief 作成
        brief = self._build_source_brief(
            context_messages=context.messages,
            conversation_excerpt=conversation_excerpt,
            attachment_summary=attachment_summary,
            file_attachments_meta=file_attachments_meta,
            research_context=research_context,
            grounding_metadata=grounding_metadata,
            system_instruction=context.system_instruction,
        )

        # 3. Presentation Structure JSON 作成
        presentation_data = self._generate_presentation_structure(
            brief=brief,
            layouts_info=layouts_info,
            has_template=has_template,
        )

        # 4. 幾何学バリデーション & 要約修復
        adjusted_slides, offsets = self._validate_and_adjust_slides(
            slides=presentation_data.slides,
            layouts_info=layouts_info,
            temp_dir=temp_dir,
        )

        # 5. スライド画像処理
        slide_images = self._prepare_slide_images(
            slides=adjusted_slides,
            temp_dir=temp_dir,
            user_images=user_images,
        )

        # 6. 物理PPTXファイル保存
        save_physical_presentation(
            presentation_data=presentation_data,
            slides=adjusted_slides,
            offsets=offsets,
            slide_images=slide_images,
            final_pptx_path=final_pptx_path,
            template_path=template_path if has_template else None,
            has_template=has_template,
            layouts_info=layouts_info,
            check_prs=check_prs,
            source_brief=brief,
        )

        # 一時作業フォルダのクリーンアップ
        try:
            shutil.rmtree(temp_dir)
        except Exception:
            pass

        return final_pptx_path


def run_pptx_agent(
    *,
    runtime: AzureRuntime,
    prompts: dict,
    context,
    messages: list[dict],
    max_output_tokens: int,
    text_placeholder,
    thought_status,
    model_id: str | None = None,
) -> tuple[str, object | None, dict[str, object]]:
    """Azure OpenAI 向けの PowerPoint 生成ディスパッチャエントリポイント。"""
    thought_status.update(label="PowerPointスライドを生成中 (Azure OpenAI)...", state="running", expanded=False)
    state_manager.add_debug_log("[AzurePPTX] PPTX generation pipeline starting...")

    # 添付画像ファイルの回収
    user_images = []
    for f_item in st.session_state.get("uploaded_file_queue", []):
        if hasattr(f_item, "type") and f_item.type.startswith("image/"):
            user_images.append({"name": f_item.name, "bytes": f_item.getvalue(), "type": f_item.type})
    for f_item in st.session_state.get("clipboard_queue", []):
        if hasattr(f_item, "type") and f_item.type.startswith("image/"):
            user_images.append({"name": f_item.name, "bytes": f_item.getvalue(), "type": f_item.type})

    enable_search = st.session_state.get("enable_google_search", False) or st.session_state.get(
        "enable_more_research", False
    )

    agent_instance = AzurePPTXAgent(runtime=runtime)
    try:
        output_pptx_path = agent_instance.generate_presentation_pipeline(
            chat_history=messages,
            context=context,
            session_id=st.session_state.get("session_id", "default_uuid"),
            user_images=user_images,
            search_enabled=enable_search,
            file_attachments_meta=context.file_attachments_meta,
        )
        full_response = (
            "Azure OpenAI により PowerPoint プレゼンテーション資料の作成が完了しました。\n\n"
            f"保存先: `{output_pptx_path}`"
        )
        thought_status.update(label="PowerPoint生成完了 (Azure OpenAI)", state="complete", expanded=False)
        report_meta = {
            "pdf_success": False,
            "llm_route": "azure_direct",
            "llm_retry_count": 0,
            "pptx_path": output_pptx_path,
        }
        text_placeholder.markdown(utils.format_latex_delimiters(full_response))
        return full_response, None, report_meta
    except Exception as e:
        tb_str = traceback.format_exc()
        full_response = f"PowerPoint資料の生成に失敗しました (Azure OpenAI): {repr(e)}"
        thought_status.update(label="PowerPoint生成失敗", state="error", expanded=True)
        state_manager.add_debug_log(f"[AzurePPTX] Generation failed: {repr(e)}\n{tb_str}", "error")
        text_placeholder.markdown(utils.format_latex_delimiters(full_response))
        raise
