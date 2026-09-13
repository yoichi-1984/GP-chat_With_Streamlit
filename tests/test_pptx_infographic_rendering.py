# tests/test_pptx_infographic_rendering.py
import os
import shutil
import tempfile
import unittest
from pptx import Presentation

from gp_chat.pptx_agent import (
    PlaceholderContent,
    PresentationDSLSchema,
    SlideNode,
    save_physical_presentation,
)


class TestPPTXInfographicRendering(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        if os.path.exists(self.temp_dir):
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_native_modern_presentation_without_template(self):
        """テンプレートが存在しない場合でも、16:9の洗練されたインフォグラフィックススライドが生成されることを検証。"""
        final_pptx = os.path.abspath(os.path.join(self.temp_dir, "native_modern_test.pptx"))

        slides = [
            # 1. KPIスライド
            SlideNode(
                slide_number=1,
                title="主要業績ハイライト",
                layout_name="blank",
                placeholders=[
                    PlaceholderContent(
                        idx=0,
                        text_content="主要指標が前年比で大幅伸長\n• 売上高: 120億円 (前年比 +15%の成長を達成)\n• 導入企業数: 1,500社 (大企業中心に拡大)\n• 顧客維持率: 98.5% (極めて高いリテンション)\n• 営業利益率: 24.0% (収益性向上)"
                    )
                ],
                visual_type="kpi",
                visual_variant="big_numbers",
                color_theme="corporate",
                accent_color_hex="#0284C7",
            ),
            # 2. プロセススライド
            SlideNode(
                slide_number=2,
                title="導入・展開プロセス",
                layout_name="blank",
                placeholders=[
                    PlaceholderContent(
                        idx=0,
                        text_content="4フェーズによる確実なシステム導入\n• Phase 1: 要件定義とPoC検証 (現状分析とスコープ確定)\n• Phase 2: クラウド基盤構築 (セキュリティ設計と環境構築)\n• Phase 3: データ移行・検証 (本番データ整合性確認)\n• Phase 4: 全社展開と運用定着 (ユーザー教育とKPI監視)"
                    )
                ],
                visual_type="process",
                visual_variant="chevron_flow",
                color_theme="corporate",
            ),
            # 3. 比較スライド (Pros & Cons)
            SlideNode(
                slide_number=3,
                title="アプローチ比較と評価",
                layout_name="blank",
                placeholders=[
                    PlaceholderContent(
                        idx=0,
                        text_content="自社開発とSaaS導入のメリット・課題\n• 自社開発のメリット: 高い柔軟性と独自要件の完全適合\n• 自社開発の課題: 初期コストと開発期間の長期化\n• SaaS導入のメリット: 即時導入可能と運用負荷の軽減\n• SaaS導入の課題: カスタマイズの制限と月額課金コスト"
                    )
                ],
                visual_type="comparison",
                visual_variant="pros_cons",
                color_theme="corporate",
            ),
            # 4. カードグリッドスライド
            SlideNode(
                slide_number=4,
                title="今後の重点施策とロードマップ",
                layout_name="blank",
                placeholders=[
                    PlaceholderContent(
                        idx=0,
                        text_content="3つの戦略的優先課題の推進\n• AI自律化の深化: エージェント機能による定型業務の完全自動化\n• セキュリティ基盤強化: ゼロトラストモデルの導入と監査ログ強化\n• エコシステム連携: 外部主要APIとのシームレスな統合の推進"
                    )
                ],
                visual_type="summary",
                visual_variant="cards_3col",
                color_theme="corporate",
            ),
        ]

        dsl = PresentationDSLSchema(
            presentation_title="2026年度事業戦略インフォグラフィックス",
            slides=slides,
        )

        save_physical_presentation(
            presentation_data=dsl,
            slides=slides,
            offsets=[0, 0, 0, 0],
            slide_images={},
            final_pptx_path=final_pptx,
            template_path=None,
            has_template=False,
            layouts_info={},
            check_prs=None,
            source_brief=None,
        )

        self.assertTrue(os.path.exists(final_pptx))
        prs = Presentation(final_pptx)
        # 16:9 ワイドスクリーン検証 (13.333 x 7.5 インチ)
        self.assertAlmostEqual(prs.slide_width.inches, 13.333, places=2)
        self.assertAlmostEqual(prs.slide_height.inches, 7.5, places=2)

        # スライド枚数（表紙 + 本文4枚 = 5枚）
        self.assertEqual(len(prs.slides), 5)

        # 表紙スライド（Slide 1）の検証
        cover_slide = prs.slides[0]
        cover_text = " ".join([sh.text for sh in cover_slide.shapes if sh.has_text_frame])
        self.assertIn("2026年度事業戦略インフォグラフィックス", cover_text)

        # 本文スライド（Slide 2〜5）の検証: 単なる箇条書きではなく複数シェイプ（カード、バッジ等）で描画されていること
        for i in range(1, 5):
            slide = prs.slides[i]
            self.assertGreaterEqual(len(slide.shapes), 4, f"Slide {i+1} should contain rich shapes, but found {len(slide.shapes)}")


if __name__ == "__main__":
    unittest.main()
