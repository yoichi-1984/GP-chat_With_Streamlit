# tests/test_azure_pptx_agent.py
import os
import unittest
from unittest.mock import MagicMock, patch

from pptx import Presentation

from gp_chat import azure_pptx_agent
from gp_chat.azure_runtime import AzureRuntime
from gp_chat.pptx_agent import (
    PlaceholderContent,
    PresentationDSLSchema,
    PresentationSourceBrief,
    SlideNode,
)


class TestAzurePPTXAgent(unittest.TestCase):
    def setUp(self):
        self.runtime = AzureRuntime(
            endpoint="https://test.openai.azure.com/",
            api_key="test-key",
            deployment="gpt-5.6-sol",
            base_url="https://test.openai.azure.com/openai/v1/",
            codex_deployment="gpt-5.3-codex",
            sol_deployment="gpt-5.6-sol",
            gpt6_deployment="gpt-6-astra",
            dalle_deployment="",
        )

    def test_agent_initialization(self):
        agent = azure_pptx_agent.AzurePPTXAgent(self.runtime)
        self.assertEqual(agent.runtime, self.runtime)

    def test_parse_source_brief_json(self):
        valid_json = """{
            "core_request": "AIの最新動向レポート",
            "audience": "エンジニア",
            "source_inventory": ["会話履歴"],
            "key_facts": ["LLMの進化", "マルチモーダル対応"],
            "evidence_notes": ["チャット内議論"],
            "visual_assets": [],
            "recommended_storyline": ["背景", "技術", "まとめ"],
            "image_policy": "不要",
            "gaps_or_uncertainties": [],
            "coverage_requirements": ["最新動向を網羅すること"],
            "source_coverage_units": ["LLMの進化", "マルチモーダル対応"],
            "references": []
        }"""
        brief = azure_pptx_agent._parse_source_brief(valid_json)
        self.assertIsInstance(brief, PresentationSourceBrief)
        self.assertEqual(brief.core_request, "AIの最新動向レポート")
        self.assertEqual(len(brief.key_facts), 2)

    def test_parse_presentation_dsl_json(self):
        valid_json = """{
            "presentation_title": "AIトレンド分析",
            "slides": [
                {
                    "slide_number": 1,
                    "title": "はじめに",
                    "layout_name": "コンテンツ 1",
                    "placeholders": [
                        {"idx": 0, "text_content": "はじめに"},
                        {"idx": 1, "text_content": "本日のアジェンダです"}
                    ],
                    "visual_type": "none",
                    "visual_variant": "auto",
                    "color_theme": "corporate",
                    "coverage_refs": ["LLMの進化"]
                }
            ]
        }"""
        schema = azure_pptx_agent._parse_presentation_dsl(valid_json)
        self.assertIsInstance(schema, PresentationDSLSchema)
        self.assertEqual(schema.presentation_title, "AIトレンド分析")
        self.assertEqual(len(schema.slides), 1)

    def test_dalle_image_generation_skipped_when_no_deployment(self):
        # dalle_deployment が空の場合は False を返し、APIを呼ばないこと
        result = azure_pptx_agent._generate_dalle_image(
            self.runtime,
            prompt="A futuristic AI workstation",
            output_path="test.png",
        )
        self.assertFalse(result)

    @patch("gp_chat.azure_responses_router._build_client")
    def test_dalle_image_generation_error_handled_safely(self, mock_build_client):
        # dalle_deployment があってもエラー時は False を返し、例外で落ちないこと
        runtime_with_dalle = AzureRuntime(
            endpoint="https://test.openai.azure.com/",
            api_key="test-key",
            deployment="gpt-5.6-sol",
            base_url="https://test.openai.azure.com/openai/v1/",
            dalle_deployment="dall-e-3",
        )
        mock_client = MagicMock()
        mock_client.images.generate.side_effect = RuntimeError("API quota exceeded")
        mock_build_client.return_value = mock_client

        result = azure_pptx_agent._generate_dalle_image(
            runtime_with_dalle,
            prompt="A futuristic AI workstation",
            output_path="test.png",
        )
        self.assertFalse(result)

    def test_physical_pptx_generation_creates_valid_file(self):
        output_dir = "tests/test_output"
        os.makedirs(output_dir, exist_ok=True)
        final_pptx = os.path.abspath(os.path.join(output_dir, "test_presentation.pptx"))
        try:
            slides = [
                SlideNode(
                    slide_number=1,
                    title="テストスライド",
                    layout_name="コンテンツ 1",
                    placeholders=[
                        PlaceholderContent(idx=0, text_content="テストタイトル"),
                        PlaceholderContent(idx=1, text_content="これはテストの本文です。\\n箇条書きの2行目。"),
                    ],
                    visual_type="none",
                    visual_variant="auto",
                    color_theme="corporate",
                )
            ]
            dsl = PresentationDSLSchema(
                presentation_title="自動テストプレゼン",
                slides=slides,
            )

            azure_pptx_agent.save_physical_presentation(
                presentation_data=dsl,
                slides=slides,
                offsets=[0],
                slide_images={},
                final_pptx_path=final_pptx,
                template_path=None,
                has_template=False,
                layouts_info={},
                check_prs=None,
                source_brief=None,
            )

            self.assertTrue(os.path.exists(final_pptx))
            # 実際に Presentation として読み込めるか検証
            prs = Presentation(final_pptx)
            self.assertGreaterEqual(len(prs.slides), 2)  # 表紙 + 本文1枚
        finally:
            if os.path.exists(final_pptx):
                os.remove(final_pptx)


if __name__ == "__main__":
    unittest.main()
